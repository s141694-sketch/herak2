"""The Word file of an approved version (task 7.3; spec 2.2 publishing, 3 module 7, 8.1 item 7; D77): the tree's
levels become heading levels in its order, the direction of each paragraph is right, and the organization's
identity is on it."""

import io

import pytest
from django.utils import timezone
from docx import Document
from docx.oxml.ns import qn

from apps.accounts.models import Organization, Role
from apps.exports import word
from apps.programs import services
from apps.programs.models import ProgramVersion
from apps.programs.tests.factories import member, program, published_template
from apps.tenancy.context import organization_context

pytestmark = pytest.mark.django_db
LEVELS = [
    {"name_ar": "وحدة", "name_en": "Unit"},
    {"name_ar": "درس", "name_en": "Lesson"},
    {"name_ar": "موضوع", "name_en": "Topic"},
]


def text(*parts, marks=None):
    return [{"type": "text", "text": part, **({"marks": marks} if marks else {})} for part in parts]


def doc(*nodes):
    return {"type": "doc", "content": list(nodes)}


def para(*content):
    return {"type": "paragraph", "content": list(content)}


@pytest.fixture
def world():
    org = Organization.objects.create(name="مركز التدريب المهني", slug="vtc", brand_colors={"primary": "#0B6E4F"})
    with organization_context(org):
        owner = member("owner@vtc.test", Role.ADMIN)
        owner.full_name = "مازن القنوبي"
        owner.save()
        version = program(
            owner, title="برنامج السلامة المهنية", template=published_template(owner, levels=LEVELS)
        ).versions.get()
        unit = services.add_node(version, title="الوحدة الأولى: السلامة", actor=owner)
        lesson = services.add_node(version, title="الدرس الأول: المعدات", actor=owner, parent=unit)
        topic = services.add_node(version, title="الموضوع: الخوذة", actor=owner, parent=lesson)
        gone = services.add_node(version, title="عقدة محذوفة", actor=owner, parent=unit)
        services.soft_delete_node(gone, actor=owner)
        second = services.add_node(version, title="الوحدة الثانية: الإسعاف", actor=owner)

        def add(node, kind, content):
            return services.add_block(version, node=node, type=kind, content=content, actor=owner)

        add(lesson, "objective", doc(para(*text("أن يحدد المتدرب معدات الوقاية."))))
        add(
            lesson,
            "objective",
            doc(para(*text("أن يطبق المتدرب "), *text("الإغلاق والتأمين", marks=[{"type": "bold"}]))),
        )
        add(
            lesson,
            "content",
            doc(
                para(*text("Personal protective equipment (PPE) protects the worker.")),
                {
                    "type": "bulletList",
                    "content": [
                        {
                            "type": "listItem",
                            "content": [
                                para(*text("الخوذة")),
                                {
                                    "type": "bulletList",
                                    "content": [{"type": "listItem", "content": [para(*text("بطانة داخلية"))]}],
                                },
                            ],
                        },
                        {"type": "listItem", "content": [para(*text("النظارة"))]},
                    ],
                },
                {
                    "type": "orderedList",
                    "attrs": {"start": 3},
                    "content": [
                        {"type": "listItem", "content": [para(*text("الخطوة الثالثة"))]},
                        {"type": "listItem", "content": [para(*text("الخطوة الرابعة"))]},
                    ],
                },
                para(
                    *text("المرجع: "),
                    *text("دليل السلامة", marks=[{"type": "link", "attrs": {"href": "https://example.com/safety"}}]),
                ),
            ),
        )
        add(topic, "activity", doc(para(*text("نشاط: فحص الخوذة"))))
        add(
            second,
            "assessment",
            doc({"type": "orderedList", "content": [{"type": "listItem", "content": [para(*text("سؤال أول"))]}]}),
        )
        ProgramVersion.objects.filter(pk=version.pk).update(approved_by=owner, approved_at=timezone.now())
        version.refresh_from_db()
    return {"org": org, "version": version, "owner": owner}


def built(world):
    with organization_context(world["org"]):
        return Document(io.BytesIO(word.build(world["version"])))


def headings(document):
    return [(p.style.name, p.text) for p in document.paragraphs if p.style.name.startswith("Heading")]


def bidi(paragraph) -> bool:
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.find(qn("w:bidi")) is not None


def test_the_trees_levels_are_the_heading_levels_in_the_trees_order(world):
    assert headings(built(world)) == [
        ("Heading 1", "الوحدة الأولى: السلامة"),
        ("Heading 2", "الدرس الأول: المعدات"),
        ("Heading 3", "الموضوع: الخوذة"),
        ("Heading 1", "الوحدة الثانية: الإسعاف"),
    ]


def test_the_document_reads_right_to_left_and_each_paragraph_takes_its_own_direction(world):
    document = built(world)
    assert document.sections[0]._sectPr.find(qn("w:bidi")) is not None
    by_text = {p.text: p for p in document.paragraphs}
    assert bidi(by_text["أن يحدد المتدرب معدات الوقاية."])
    assert not bidi(by_text["Personal protective equipment (PPE) protects the worker."])
    assert all(bidi(p) for p in document.paragraphs if p.style.name.startswith("Heading"))


def test_each_kind_of_block_is_introduced_once_and_formatting_is_kept(world):
    document = built(world)
    texts = [p.text for p in document.paragraphs]
    objectives = texts.index("الأهداف")
    assert texts[objectives + 1 : objectives + 3] == [
        "أن يحدد المتدرب معدات الوقاية.",
        "أن يطبق المتدرب الإغلاق والتأمين",
    ]
    assert texts.count("الأهداف") == 1 and "المحتوى" in texts and "الأنشطة" in texts and "التقويم" in texts
    second = next(p for p in document.paragraphs if p.text == "أن يطبق المتدرب الإغلاق والتأمين")
    assert [run.bold for run in second.runs] == [None, True]


def test_lists_keep_their_nesting_and_numbered_ones_their_start(world):
    document = built(world)
    styles = {p.text: p.style.name for p in document.paragraphs}
    assert styles["الخوذة"] == "List Bullet" and styles["بطانة داخلية"] == "List Bullet 2"
    third = next(p for p in document.paragraphs if p.text == "الخطوة الثالثة")
    num_id = int(third._p.pPr.numPr.numId.val)
    numbering = document.part.numbering_part.element
    num = next(n for n in numbering.findall(qn("w:num")) if int(n.get(qn("w:numId"))) == num_id)
    assert num.find(qn("w:lvlOverride")).find(qn("w:startOverride")).get(qn("w:val")) == "3"
    question = next(p for p in document.paragraphs if p.text == "سؤال أول")
    assert int(question._p.pPr.numPr.numId.val) != num_id  # every numbered list starts again


def test_links_stay_links(world):
    document = built(world)
    links = [rel.target_ref for rel in document.part.rels.values() if rel.reltype.endswith("/hyperlink")]
    assert links == ["https://example.com/safety"]
    assert any("دليل السلامة" in p.text for p in document.paragraphs)


def test_the_cover_carries_the_organization_the_program_and_its_approval(world):
    document = built(world)
    cover = "\n".join(p.text for p in document.paragraphs[:8])
    assert "مركز التدريب المهني" in cover and "برنامج السلامة المهنية" in cover
    assert "النسخة 1" in cover and "مازن القنوبي" in cover
    assert document.core_properties.title == "برنامج السلامة المهنية"
    assert str(document.styles["Heading 1"].font.color.rgb) == "0B6E4F"


def test_deleted_parts_of_the_tree_are_left_out(world):
    assert "عقدة محذوفة" not in [p.text for p in built(world).paragraphs]


def test_bold_and_sizes_also_apply_to_arabic_letters(world):
    """Word and LibreOffice style Arabic with the complex-script properties (bCs, iCs, szCs), not b, i and sz."""
    document = built(world)
    bold = next(p for p in document.paragraphs if p.text == "أن يطبق المتدرب الإغلاق والتأمين").runs[1]
    assert bold._r.rPr.find(qn("w:bCs")) is not None
    title = next(p for p in document.paragraphs if p.text == "برنامج السلامة المهنية").runs[0]
    assert title._r.rPr.find(qn("w:szCs")).get(qn("w:val")) == "52"
    introduction = next(p for p in document.paragraphs if p.text == "الأهداف").runs[0]
    assert introduction._r.rPr.find(qn("w:bCs")) is not None


def test_the_organizations_logo_is_on_the_cover(world):
    from apps.accounts.tests.test_identity import PNG

    with organization_context(world["org"]):
        document = Document(io.BytesIO(word.build(world["version"], logo=PNG)))
    assert len(document.inline_shapes) == 1
