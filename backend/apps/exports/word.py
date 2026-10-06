"""The Word file of a version (task 7.3; spec 2.2 publishing, 3 module 7; D77), written with python-docx from the
version's rows: a cover with the organization's identity, then the tree, its levels as heading levels in its
order, each node's blocks under it. The document reads right to left; each paragraph takes the direction of its
first strong letter, so English passages stay left to right."""

import io
import re
import unicodedata
from collections import defaultdict
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from apps.programs.models import ProgramVersion

from .labels import label

FONT = "Arial"  # in Word everywhere, and with Arabic letters; LibreOffice falls back to a font that has them
HEX = re.compile(r"^#?([0-9A-Fa-f]{6})$")
BULLETS = ["List Bullet", "List Bullet 2", "List Bullet 3"]
NUMBERS = ["List Number", "List Number 2", "List Number 3"]
TIMEZONE = ZoneInfo("Asia/Muscat")


def build(version: ProgramVersion, *, logo: bytes | None = None) -> bytes:
    program = version.program
    organization = program.organization
    color = _brand_color(organization.brand_colors)
    document = Document()
    _setup(document, color)
    section = document.sections[0]
    section._sectPr.append(_element("w:bidi"))
    _cover(document, version, logo=logo)
    body = document.add_section(WD_SECTION.NEW_PAGE)
    _header_and_footer(body, organization.name, program.title)
    _tree(document, version)
    document.core_properties.title = program.title
    document.core_properties.author = organization.name
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# --- The document's look -----------------------------------------------------------------------------------


def _brand_color(colors) -> RGBColor | None:
    value = colors.get("primary") if isinstance(colors, dict) else None
    match = HEX.match(value) if isinstance(value, str) else None
    return RGBColor.from_string(match.group(1).upper()) if match else None


def _setup(document, color: RGBColor | None) -> None:
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    _fonts(normal.element.get_or_add_rPr(), size=24)
    rpr = normal.element.get_or_add_rPr()
    lang = _element("w:lang", {"w:val": "en-US", "w:bidi": "ar-SA"})
    rpr.append(lang)
    for level in range(1, 6):
        style = document.styles[f"Heading {level}"]
        style.font.name = FONT
        _fonts(style.element.get_or_add_rPr())
        if color is not None:
            style.font.color.rgb = color
    for margin in ("left_margin", "right_margin"):
        setattr(document.sections[0], margin, Cm(2.2))


def _fonts(rpr, size: int | None = None) -> None:
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = _element("w:rFonts")
        rpr.insert(0, fonts)
    for slot in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(slot), FONT)
    if size is not None:
        rpr.append(_element("w:szCs", {"w:val": str(size)}))


def _format(run, *, bold: bool = False, italic: bool = False, size: Pt | None = None) -> None:
    """Bold, italic and size for every script: Arabic takes the complex-script properties (bCs, iCs, szCs)."""
    rpr = run._r.get_or_add_rPr()
    if bold:
        run.bold = True
        rpr.append(_element("w:bCs"))
    if italic:
        run.italic = True
        rpr.append(_element("w:iCs"))
    if size is not None:
        run.font.size = size
        rpr.append(_element("w:szCs", {"w:val": str(int(size.pt * 2))}))


def _element(tag: str, attrs: dict | None = None):
    element = OxmlElement(tag)
    for key, value in (attrs or {}).items():
        element.set(qn(key), value)
    return element


# --- Direction -------------------------------------------------------------------------------------------------


def right_to_left(text: str) -> bool:
    """The direction of the first strong letter; right to left when there is none (the document's own)."""
    for char in text:
        kind = unicodedata.bidirectional(char)
        if kind in ("R", "AL"):
            return True
        if kind == "L":
            return False
    return True


def _direct(paragraph, text: str) -> None:
    if right_to_left(text):
        paragraph._p.get_or_add_pPr().insert(0, _element("w:bidi"))
        for run in paragraph.runs:
            run._r.get_or_add_rPr().append(_element("w:rtl"))


# --- Cover, header and footer ------------------------------------------------------------------------------


def _cover(document, version: ProgramVersion, *, logo: bytes | None) -> None:
    program = version.program
    if logo:
        picture = document.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.add_run().add_picture(io.BytesIO(logo), height=Cm(3))
    lines = [(program.organization.name, Pt(16), True), (program.title, Pt(26), True)]
    if program.target_role:
        lines.append((label("target_role", role=program.target_role), Pt(13), False))
    lines.append((label("version", number=version.number), Pt(13), False))
    if version.approved_at and version.approved_by:
        when = version.approved_at.astimezone(TIMEZONE).strftime("%Y-%m-%d")
        lines.append(
            (
                label("approved", name=version.approved_by.full_name or version.approved_by.email, date=when),
                Pt(12),
                False,
            )
        )
    for text, size, bold in lines:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _format(paragraph.add_run(text), bold=bold, size=size)
        _direct(paragraph, text)


def _header_and_footer(section, organization: str, title: str) -> None:
    section.different_first_page_header_footer = False
    header = section.header
    header.is_linked_to_previous = False
    head = header.paragraphs[0]
    _format(head.add_run(f"{organization} — {title}"), size=Pt(9))
    _direct(head, head.text)
    footer = section.footer
    footer.is_linked_to_previous = False
    foot = footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot.add_run(f"{label('page')} ")
    field = _element("w:fldSimple", {"w:instr": "PAGE"})
    field.append(OxmlElement("w:r"))
    foot._p.append(field)
    _direct(foot, label("page"))


# --- The tree --------------------------------------------------------------------------------------------------


def _tree(document, version: ProgramVersion) -> None:
    nodes = list(version.nodes.filter(deleted=False).order_by("level", "order", "pk"))
    blocks = defaultdict(list)
    for block in version.blocks.filter(deleted=False, node__deleted=False).order_by("order", "pk"):
        blocks[block.node_id].append(block)
    children = defaultdict(list)
    alive = {node.pk for node in nodes}
    for node in nodes:
        if node.parent_id is None or node.parent_id in alive:
            children[node.parent_id].append(node)
    writer = _Writer(document)

    def visit(node, depth: int) -> None:
        heading = document.add_heading(node.title, level=min(depth, 5))
        _direct(heading, node.title)
        kind = None
        for block in blocks.get(node.pk, []):
            if block.type != kind:  # each kind of block is introduced once in a row
                kind = block.type
                introduction = document.add_paragraph()
                _format(introduction.add_run(label(kind)), bold=True)
                _direct(introduction, introduction.text)
            writer.content(block.content)
        for child in children.get(node.pk, []):
            visit(child, depth + 1)

    for root in children.get(None, []):
        visit(root, 1)


class _Writer:
    """Writes a block's rich content (the schema of apps.programs.content) as Word paragraphs."""

    def __init__(self, document):
        self.document = document
        self.numbering = document.part.numbering_part.element

    def content(self, doc: dict) -> None:
        for node in doc.get("content", []):
            self.block(node, depth=0)

    def block(self, node: dict, *, depth: int, style: str | None = None, number: int | None = None) -> None:
        kind = node.get("type")
        if kind == "paragraph":
            self.paragraph(node.get("content", []), style=style, number=number)
        elif kind == "heading":
            paragraph = self.paragraph(node.get("content", []))
            for run in paragraph.runs:
                _format(run, bold=True, size=Pt(14))
        elif kind == "bulletList":
            for item in node.get("content", []):
                self.item(item, depth=depth, style=BULLETS[min(depth, 2)])
        elif kind == "orderedList":
            number = self.restarted(NUMBERS[min(depth, 2)], (node.get("attrs") or {}).get("start") or 1)
            for item in node.get("content", []):
                self.item(item, depth=depth, style=NUMBERS[min(depth, 2)], number=number)
        elif kind == "blockquote":
            for child in node.get("content", []):
                self.block(child, depth=depth, style="Quote")
        elif kind == "codeBlock":
            paragraph = self.document.add_paragraph()
            run = paragraph.add_run("".join(part.get("text", "") for part in node.get("content", [])))
            run.font.name = "Courier New"
        elif kind == "horizontalRule":
            paragraph = self.document.add_paragraph()
            border = _element("w:pBdr")
            border.append(_element("w:bottom", {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": "auto"}))
            paragraph._p.get_or_add_pPr().append(border)

    def item(self, item: dict, *, depth: int, style: str, number: int | None = None) -> None:
        first = True
        for child in item.get("content", []):
            if child.get("type") in ("bulletList", "orderedList"):
                self.block(child, depth=depth + 1)
            else:
                # The item's first paragraph carries its bullet or number; the rest continue it.
                self.block(child, depth=depth, style=style if first else None, number=number if first else None)
                first = False

    def restarted(self, style: str, start: int) -> int:
        """A numbering of its own for this list, from ``start``: Word would otherwise go on from the last one."""
        base = self.document.styles[style].element.pPr.numPr.numId.val
        abstract = next(
            n.find(qn("w:abstractNumId")).get(qn("w:val"))
            for n in self.numbering.findall(qn("w:num"))
            if int(n.get(qn("w:numId"))) == int(base)
        )
        number = max(int(n.get(qn("w:numId"))) for n in self.numbering.findall(qn("w:num"))) + 1
        num = _element("w:num", {"w:numId": str(number)})
        num.append(_element("w:abstractNumId", {"w:val": abstract}))
        override = _element("w:lvlOverride", {"w:ilvl": "0"})
        override.append(_element("w:startOverride", {"w:val": str(int(start))}))
        num.append(override)
        self.numbering.append(num)
        return number

    def paragraph(self, inline: list, *, style: str | None = None, number: int | None = None):
        paragraph = self.document.add_paragraph(style=style)
        if number is not None:
            numpr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
            numpr.get_or_add_ilvl().val = 0
            numpr.get_or_add_numId().val = number
        for part in inline:
            if part.get("type") == "hardBreak":
                (paragraph.runs[-1] if paragraph.runs else paragraph.add_run()).add_break(WD_BREAK.LINE)
            elif part.get("type") == "text":
                self.text(paragraph, part)
        _direct(paragraph, paragraph.text)
        return paragraph

    def text(self, paragraph, part: dict) -> None:
        marks = {mark.get("type"): mark.get("attrs") or {} for mark in part.get("marks") or []}
        if "link" in marks:
            self.link(paragraph, part["text"], marks["link"].get("href", ""), marks)
            return
        run = paragraph.add_run(part.get("text", ""))
        self.style(run, marks)

    @staticmethod
    def style(run, marks: dict) -> None:
        _format(run, bold="bold" in marks, italic="italic" in marks)
        if "underline" in marks:
            run.underline = True
        if "strike" in marks:
            run.font.strike = True
        if "code" in marks:
            run.font.name = "Courier New"

    def link(self, paragraph, text: str, href: str, marks: dict) -> None:
        relation = paragraph.part.relate_to(href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = _element("w:hyperlink", {"r:id": relation})
        run = paragraph.add_run(text)
        self.style(run, marks)
        run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
        run.underline = True
        paragraph._p.remove(run._r)
        hyperlink.append(run._r)
        paragraph._p.append(hyperlink)
