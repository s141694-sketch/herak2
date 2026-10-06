"""Text from an uploaded curriculum (task 7.1, spec 2.2 production, 7.6, 7.7): PDF with pypdf, Word with
python-docx, and plain text; the type is read from the content, and what cannot be read is refused plainly."""

import io
import time
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, StreamObject

from apps.files import extraction

FIXTURES = Path(__file__).parent / "fixtures"


def word(*paragraphs, table=None) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, text in enumerate(row):
                grid.cell(r, c).text = text
    document.add_paragraph("الخاتمة")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_a_word_file_gives_its_paragraphs_and_tables_in_reading_order():
    data = word("الوحدة الأولى", "الأهداف", table=[["الهدف", "المستوى"], ["أن يحدد المتدرب المعدات", "تطبيق"]])
    extracted = extraction.extract(data, "منهج.docx")
    assert extracted.kind == "docx"
    lines = [line for line in extracted.text.split("\n") if line]
    assert lines == ["الوحدة الأولى", "الأهداف", "الهدف", "المستوى", "أن يحدد المتدرب المعدات", "تطبيق", "الخاتمة"]


def test_an_arabic_pdf_gives_its_text_in_reading_order():
    extracted = extraction.extract((FIXTURES / "curriculum.pdf").read_bytes(), "curriculum.pdf")
    assert extracted.kind == "pdf"
    text = extracted.text
    assert text.index("الوحدة الأولى: السلامة في الورشة") < text.index("أن يحدد المتدرب معدات الوقاية الشخصية")


def test_plain_text_in_utf8_with_or_without_a_byte_order_mark():
    assert extraction.extract("نص عربي".encode(), "a.txt").text == "نص عربي"
    assert extraction.extract("﻿نص عربي".encode(), "a.txt").text == "نص عربي"


def test_the_type_comes_from_the_content_not_the_name():
    assert extraction.extract(word("نص"), "renamed.pdf").kind == "docx"
    assert extraction.extract((FIXTURES / "curriculum.pdf").read_bytes(), "renamed.docx").kind == "pdf"


@pytest.mark.parametrize(
    "data,name,code",
    [
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64, "old.doc", "file_type_unsupported"),
        (b"\x89PNG\r\n\x1a\n" + b"\0" * 64, "scan.png", "file_type_unsupported"),
        (b"%PDF-1.7 broken", "broken.pdf", "file_unreadable"),
        ("نص".encode("cp1256"), "legacy.txt", "file_unreadable"),
    ],
)
def test_what_cannot_be_read_is_refused_with_a_reason(data, name, code):
    with pytest.raises(extraction.Refused) as refused:
        extraction.extract(data, name)
    assert refused.value.code == code


def test_a_pdf_locked_with_a_password_is_refused():
    with pytest.raises(extraction.Refused) as refused:
        extraction.extract((FIXTURES / "encrypted.pdf").read_bytes(), "encrypted.pdf")
    assert refused.value.code == "file_encrypted"


def test_a_file_over_the_size_limit_is_refused_before_it_is_read(settings):
    settings.IMPORT_MAX_FILE_BYTES = 1000
    with pytest.raises(extraction.Refused) as refused:
        extraction.extract(b"x" * 1001, "a.txt")
    assert refused.value.code == "import_file_too_large"


def test_a_pdf_with_too_many_pages_is_refused(settings):
    settings.IMPORT_MAX_PAGES = 0
    with pytest.raises(extraction.Refused) as refused:
        extraction.extract((FIXTURES / "curriculum.pdf").read_bytes(), "curriculum.pdf")
    assert refused.value.code == "file_too_many_pages"


def test_a_word_file_that_would_unpack_to_much_more_than_its_size_is_refused(settings):
    settings.IMPORT_MAX_UNPACKED_BYTES = 10_000
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(word("نص"))) as original, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as bomb:
        for item in original.infolist():
            bomb.writestr(item, original.read(item.filename))
        bomb.writestr("word/media/padding.bin", b"\0" * 50_000)
    with pytest.raises(extraction.Refused) as refused:
        extraction.extract(out.getvalue(), "big.docx")
    assert refused.value.code == "import_file_too_large"


def test_a_pdf_without_text_gives_no_text():
    """A scanned PDF has pages but no text: the caller then says to paste the text instead."""
    assert extraction.extract((FIXTURES / "scanned.pdf").read_bytes(), "scanned.pdf").text.strip() == ""


# From the phase 7 review: what a file can make the server do is bounded (D80).


def test_a_word_cell_spanning_millions_of_columns_is_read_at_once():
    """python-docx's row.cells makes one entry per spanned column: 150 million took 15 s and 1.5 GB."""
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "خلية واسعة"
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:gridSpan {nsdecls("w")} w:val="150000000"/>'))
    out = io.BytesIO()
    document.save(out)
    started = time.monotonic()
    assert "خلية واسعة" in extraction.extract(out.getvalue(), "wide.docx").text
    assert time.monotonic() - started < 5


def test_a_table_python_docx_cannot_lay_out_is_still_read():
    """A cell continuing a vertical merge with nothing above it made python-docx raise, and the API answer 500."""
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "خلية"
    cell._tc.get_or_add_tcPr().append(parse_xml(f"<w:vMerge {nsdecls('w')}/>"))
    out = io.BytesIO()
    document.save(out)
    assert "خلية" in extraction.extract(out.getvalue(), "merged.docx").text


def test_a_damaged_word_file_is_refused_as_unreadable():
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(word("نص"))) as original, zipfile.ZipFile(out, "w") as damaged:
        for item in original.infolist():
            data = original.read(item.filename)
            damaged.writestr(item, data[: len(data) // 2] if item.filename == "word/document.xml" else data)
    with pytest.raises(extraction.Refused) as refused:
        extraction.extract(out.getvalue(), "damaged.docx")
    assert refused.value.code == "file_unreadable"


def heavy_pdf(pages: int) -> bytes:
    """A small PDF whose pages share one long content stream: slow to read, though only a few kilobytes."""
    writer = PdfWriter()
    content = StreamObject()
    content.set_data(b"0 0 m 1 1 l S\n" * 80_000)
    content = writer._add_object(content.flate_encode())
    font = {NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1")}
    font[NameObject("/BaseFont")] = NameObject("/Helvetica")
    fonts = DictionaryObject({NameObject("/F1"): writer._add_object(DictionaryObject(font))})
    resources = DictionaryObject({NameObject("/Font"): fonts})
    for _ in range(pages):
        page = writer.add_blank_page(100, 100)
        page[NameObject("/Contents")] = content
        page[NameObject("/Resources")] = resources
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def test_reading_a_file_stops_at_the_time_allowed(settings):
    settings.IMPORT_EXTRACT_SECONDS = 1
    data = heavy_pdf(40)
    assert len(data) < 20_000
    started = time.monotonic()
    with pytest.raises(extraction.Refused) as refused:
        extraction.extract(data, "heavy.pdf")
    assert refused.value.code == "file_unreadable"
    assert time.monotonic() - started < 4


def test_text_longer_than_an_import_takes_is_refused_as_too_long():
    from apps.agents.importing import MAX_TEXT_CHARS

    with pytest.raises(extraction.Refused) as refused:
        extraction.extract(("ن" * (MAX_TEXT_CHARS + 1)).encode(), "long.txt")
    assert refused.value.code == "file_text_too_long"


def test_a_word_file_made_by_libreoffice_reads_in_order():
    """Plan 7.1 asks for files from Word or LibreOffice, not only python-docx (phase 7 review). This one was saved by
    LibreOffice from an Arabic page: headings, a numbered list and a table."""
    data = (FIXTURES / "curriculum-libreoffice.docx").read_bytes()
    assert b"LibreOffice" in zipfile.ZipFile(io.BytesIO(data)).read("docProps/app.xml")
    lines = [line for line in extraction.extract(data, "منهج.docx").text.split("\n") if line.strip()]
    assert lines[:3] == ["الوحدة الأولى: السلامة في الورشة", "الدرس الأول: معدات الوقاية الشخصية", "الأهداف"]
    assert (
        lines.index("فحص الخوذة")
        > lines.index("النشاط")
        > lines.index("أن يطبق المتدرب إجراءات الإغلاق والتأمين قبل بدء أعمال الصيانة.")
    )
    assert lines[-1] == "الخاتمة"
